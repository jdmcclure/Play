from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from assigner import group_by_person

HEADER_FILL = PatternFill("solid", fgColor="1E4D2B")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=12)
STRIPE_FILL = PatternFill("solid", fgColor="EEF4EF")
THIN_BORDER = Border(bottom=Side(style="thin", color="D5DBD6"))
MAX_COLUMN_WIDTH = 70


def _write_sheet(sheet, headers, rows):
    sheet.append(headers)
    for row in rows:
        sheet.append(row)

    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 24

    for row_number, row in enumerate(sheet.iter_rows(min_row=2), start=2):
        line_count = 1
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = THIN_BORDER
            if row_number % 2 == 1:
                cell.fill = STRIPE_FILL
            line_count = max(line_count, str(cell.value or "").count("\n") + 1)
        sheet.row_dimensions[row_number].height = max(18, 15 * line_count)

    # Size each column to its longest line of text
    for column_number, header in enumerate(headers, start=1):
        longest = len(str(header))
        for row in rows:
            for line in str(row[column_number - 1]).split("\n"):
                longest = max(longest, len(line))
        sheet.column_dimensions[get_column_letter(column_number)].width = min(longest + 4, MAX_COLUMN_WIDTH)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def export_to_excel(assignments, names, file_path):
    """Write the assignments to an Excel workbook with a 'By Measure' and a 'By Person' sheet."""
    workbook = Workbook()

    by_measure = workbook.active
    by_measure.title = "By Measure"
    _write_sheet(by_measure, ["Ballot Measure", "Assigned To"], [list(pair) for pair in assignments])

    by_person = workbook.create_sheet("By Person")
    person_rows = [
        [name, len(measures), "\n".join(measures)]
        for name, measures in group_by_person(assignments, names).items()
    ]
    _write_sheet(by_person, ["Name", "Count", "Ballot Measures"], person_rows)

    workbook.properties.title = "Ballot Measure Assignments"
    workbook.properties.created = datetime.now()
    workbook.save(file_path)
